"""
Exportação do Cronograma (atividades) em planilha Excel — botão "Exportar
Cronograma (Excel)" na tela de Cronograma, ao lado de Importar/Versões.

Diferente da importação (que lê uma planilha e grava atividades) e da
exportação de "Minhas atividades" (apontamento de horas, escopo todos os
projetos), esta é uma exportação simples e 100% determinística: pega as
atividades do projeto selecionado — já filtradas no backend por Etapa,
Frente de trabalho e/ou período previsto, todos opcionais — e devolve numa
única aba, achatadas (sem agrupamento), na mesma ordem/lógica de colunas já
usada na tabela do Cronograma no front-end (ver renderAtividadesTable em
frontend/index.html), com algumas colunas a mais que só cabem numa planilha
(datas reais, observações, lista de responsáveis).

`gerar_planilha_bytes(atividades, dependencias_por_atividade=None)` recebe a
lista de dicts já filtrada (cada item no formato devolvido por
ATIVIDADE_SELECT, ver main.py) e devolve os bytes do .xlsx — nenhuma consulta
ao banco acontece aqui, só formatação. `dependencias_por_atividade` é opcional:
um dict {atividade_id: [{predecessora_codigo_wbs, tipo, lag_horas}, ...]},
usado só para preencher a coluna "Depende de" (ver abaixo).

45ª rodada: pedido do usuário verbatim: *"A função extração deve extrair
todos os campos, assim, posso fazer manutenção e carregar de novo."* — até
então a planilha exportada não cobria todos os campos que a tela "Editar em
massa" já permite editar (ver cronograma_edicao_lote.py, 43ª/44ª rodada):
faltavam Tipo de atividade elementar, Item do TR, Depende de, Descrição e
Objetivo. Os cinco foram adicionados como novas colunas, antes de
"Observações" (que continua sendo a última, por ser o campo de texto livre
mais comprido). Isso fecha o ciclo "exportar → editar no Excel → reimportar"
pra QUALQUER campo de cadastro da atividade — o modo "flat" da importação
(ver cronograma_import.py) já reconhece e regrava todas essas colunas de
volta, com a mesma cautela de sempre: célula em branco NUNCA apaga o que já
está gravado (pra limpar um campo, ou remover TODAS as predecessoras de uma
atividade, use a tela "Editar em massa").

- **Tipo de atividade elementar**: nome do tipo (cadastro global) — casado
  por nome na reimportação, mesmo texto mostrado na tela.
- **Item do TR**: "{Código} — {Título}" (mesmo formato usado na tela) —
  casado pelo Código (a parte antes do travessão) na reimportação.
- **Depende de**: cada predecessora como "{Código} (TIPO)" ou "{Código}
  (TIPO, lag Xh)" quando há lag, várias separadas por "; " — ex.:
  "A.1 (FS); A.3 (SS, lag 8h)". Predecessora sem Código (não dá pra
  referenciar por texto) fica de fora da célula. Na reimportação, o texto da
  célula é tratado como a lista COMPLETA de predecessoras daquela linha —
  diferente dos outros campos, aqui uma célula preenchida também REMOVE uma
  predecessora que estava gravada mas não está mais na lista (é a forma da
  planilha de description "eu quero que a lista seja exatamente esta"); só
  uma célula vazia não mexe em nada, mesma regra de sempre.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

COR_CABECALHO_HEX = "1A3D5C"

CABECALHO = [
    "Código", "Atividade", "Etapa", "Frente", "Status", "Prioridade",
    "Início prev.", "Fim prev.", "Esforço prev. (h)",
    "Início real", "Fim real", "Horas realizadas (h)", "% concluído",
    "Atrasada", "Master", "Entregável", "Responsáveis",
    "Tipo de atividade elementar", "Item do TR", "Depende de",
    "Descrição", "Objetivo", "Observações",
]
LARGURAS = [12, 36, 22, 20, 16, 11, 13, 13, 14, 13, 13, 15, 12, 10, 8, 10, 28,
            24, 26, 30, 34, 30, 34]


def _estilo_cabecalho(ws, linha, n_colunas):
    for col in range(1, n_colunas + 1):
        cel = ws.cell(row=linha, column=col)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill(start_color=COR_CABECALHO_HEX, end_color=COR_CABECALHO_HEX, fill_type="solid")
        cel.alignment = Alignment(vertical="center")


def _autofit(ws, larguras):
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura


def _fmt_data_br(d):
    if not d:
        return None
    # db.fetch_all devolve date/datetime já como string ISO ("YYYY-MM-DD") ou
    # objeto date, dependendo do driver — cobre os dois casos sem depender de
    # nenhum import extra (mesmo padrão defensivo usado no resto do projeto).
    s = d if isinstance(d, str) else d.isoformat()
    partes = s.split("-")
    return f"{partes[2]}/{partes[1]}/{partes[0]}" if len(partes) == 3 else s


def _responsaveis_txt(responsaveis):
    if not responsaveis:
        return ""
    return ", ".join(f"{r['nome']}" + (f" ({r['tipo_vinculo']})" if r.get("tipo_vinculo") else "") for r in responsaveis)


def _item_tr_txt(a):
    codigo = a.get("requisito_tr_codigo")
    if not codigo:
        return ""
    titulo = a.get("requisito_tr_titulo")
    return f"{codigo} — {titulo}" if titulo else codigo


def _lag_txt(lag_horas):
    try:
        lag = float(lag_horas) if lag_horas not in (None, "") else 0.0
    except (TypeError, ValueError):
        lag = 0.0
    # sem casas decimais quando é um número redondo (ex.: 8, não 8.0)
    return f"{lag:g}"


def _depende_de_txt(deps):
    if not deps:
        return ""
    itens = []
    for d in deps:
        codigo = d.get("predecessora_codigo_wbs")
        if not codigo:
            # sem Código não dá pra referenciar por texto numa reimportação —
            # mesma limitação já existente pra qualquer outra coluna que casa
            # atividades pelo Código (ver módulo cronograma_import.py).
            continue
        tipo = d.get("tipo") or "FS"
        lag_horas = d.get("lag_horas")
        lag = float(lag_horas) if lag_horas not in (None, "") else 0.0
        itens.append(f"{codigo} ({tipo}, lag {_lag_txt(lag)}h)" if lag else f"{codigo} ({tipo})")
    return "; ".join(itens)


def gerar_planilha_bytes(atividades, dependencias_por_atividade=None):
    dependencias_por_atividade = dependencias_por_atividade or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Cronograma"
    ws.append(CABECALHO)
    _estilo_cabecalho(ws, 1, len(CABECALHO))

    for a in atividades:
        ws.append([
            a.get("codigo_wbs") or "",
            a.get("nome") or "",
            (f"{a['etapa_numero']} — {a['etapa_nome']}" if a.get("etapa_nome") else ""),
            a.get("frente_nome") or "",
            a.get("status") or "",
            a.get("prioridade") or "",
            _fmt_data_br(a.get("dtini_prev")),
            _fmt_data_br(a.get("dtfim_prev")),
            float(a["prazo_horas"]) if a.get("prazo_horas") is not None else None,
            _fmt_data_br(a.get("dtini_real")),
            _fmt_data_br(a.get("dtfim_real")),
            float(a["horas_realizadas"]) if a.get("horas_realizadas") is not None else None,
            a.get("percentual_concluido"),
            "Sim" if a.get("atrasada") else "Não",
            "Sim" if a.get("eh_atividade_master") else "",
            "Sim" if a.get("eh_entregavel") else "",
            _responsaveis_txt(a.get("responsaveis")),
            a.get("tipo_nome") or "",
            _item_tr_txt(a),
            _depende_de_txt(dependencias_por_atividade.get(a.get("id"))),
            a.get("descricao") or "",
            a.get("objetivo") or "",
            a.get("observacoes") or "",
        ])

    _autofit(ws, LARGURAS)
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
