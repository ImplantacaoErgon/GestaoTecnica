"""
Exportação em Excel do "Quadro de Carga por Ciclo" — Migração de Dados
(78ª/79ª/80ª rodada).

O quadro em si (filtro de ciclo, linhas por item de migração, total geral)
já é montado e calculado inteiramente no backend por
`quadro_ciclo_migracao()` em app/main.py — este módulo só formata os
MESMOS dados (nenhum recálculo aqui) numa planilha de uma aba só, no
mesmo espírito de app/relatorio_atividades.py: cabeçalho colorido, colunas
com largura fixada na unha (autofit manual, sem depender de biblioteca
extra) e uma linha de total geral em negrito ao final.
"""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

COR_CABECALHO_HEX = "1A3D5C"
COR_TOTAL_HEX = "E8EDF2"

CABECALHO = [
    "Destino/Sistema", "Sistema de origem", "Número do Ciclo", "Data do ciclo",
    "Total a carregar", "Total carregado", "Total rejeitados",
    "% carregado", "% rejeitado",
]
LARGURAS = [28, 22, 14, 14, 16, 16, 16, 13, 13]
COL_PCT = (8, 9)  # % carregado, % rejeitado — 1-based, ver CABECALHO acima
COL_DATA = 4


def _pct(v):
    """Recebe o percentual já calculado (0-100, ou None) e devolve como
    fração 0-1 pra célula usar formato de porcentagem nativo do Excel (em
    vez de guardar "83.33%" como texto) — assim a planilha permite somar/
    fazer média/gráfico com a coluna normalmente."""
    return (float(v) / 100.0) if v is not None else None


def _data(v):
    """80ª rodada: `data_execucao` chega do banco como string "AAAA-MM-DD"
    (serializada em JSON) — converte pra `date` de verdade, não texto, pra
    célula usar o formato de data nativo do Excel (ordena/filtra como data,
    não como string)."""
    if v is None:
        return None
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def gerar_planilha_bytes(numero_ciclo, linhas, totais):
    """`linhas`: mesma lista devolvida por GET /itens-migracao/quadro-ciclo
    (campos nome_tabela_destino/nome_tabela_legado, sistema_origem,
    numero_ciclo, data_execucao, total_a_carregar, total_carregado,
    total_rejeitados, percentual_carregado, percentual_rejeicao). `totais`:
    dict com total_a_carregar/total_carregado/total_rejeitados/
    percentual_carregado/percentual_rejeicao já somados (mesmo cálculo do
    total geral exibido na tela — não é recalculado aqui de novo)."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Ciclo {numero_ciclo}" if numero_ciclo is not None else "Quadro"

    ws.append(CABECALHO)
    for col in range(1, len(CABECALHO) + 1):
        cel = ws.cell(row=1, column=col)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill(start_color=COR_CABECALHO_HEX, end_color=COR_CABECALHO_HEX, fill_type="solid")
        cel.alignment = Alignment(vertical="center")

    linha = 2
    for l in linhas:
        ws.append([
            l.get("nome_tabela_destino") or l.get("nome_tabela_legado") or "",
            l.get("sistema_origem") or "",
            l.get("numero_ciclo"),
            _data(l.get("data_execucao")),
            l.get("total_a_carregar") or 0,
            l.get("total_carregado") or 0,
            l.get("total_rejeitados") or 0,
            _pct(l.get("percentual_carregado")),
            _pct(l.get("percentual_rejeicao")),
        ])
        ws.cell(row=linha, column=COL_DATA).number_format = "dd/mm/yyyy"
        for col in COL_PCT:
            ws.cell(row=linha, column=col).number_format = "0.00%"
        linha += 1

    if linhas:
        ws.append([
            "Total geral", "", "", "",
            totais.get("total_a_carregar") or 0,
            totais.get("total_carregado") or 0,
            totais.get("total_rejeitados") or 0,
            _pct(totais.get("percentual_carregado")),
            _pct(totais.get("percentual_rejeicao")),
        ])
        for col in range(1, len(CABECALHO) + 1):
            cel = ws.cell(row=linha, column=col)
            cel.font = Font(bold=True)
            cel.fill = PatternFill(start_color=COR_TOTAL_HEX, end_color=COR_TOTAL_HEX, fill_type="solid")
        for col in COL_PCT:
            ws.cell(row=linha, column=col).number_format = "0.00%"

    for i, largura in enumerate(LARGURAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
