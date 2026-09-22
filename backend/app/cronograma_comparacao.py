"""
Comparação entre uma versão salva do cronograma e o cronograma AO VIVO —
botão "📊 Comparar" na lista de "Versões do cronograma" (ver
cronograma_versoes.py). 46ª rodada, pedido do usuário verbatim: *"criar uma
função que permita a geração de uma planilha com a comparação de cronograma
atual com alguma das versões salvas. A planilha deverá conter as atividades
da versão salva acrescido das atividades do Cronograma atual — se possível
com a cor de azul bem clara — e as diferenças de dados da mesma atividade
apresentar as informações do cronograma atual na cor amarela"*.

Regras (implementação do pedido acima):
- Uma linha por atividade, casada pelo **id** (não pelo Código WBS — mais
  robusto, porque o id nunca muda mesmo que o Código seja editado depois da
  versão ter sido gerada; é o mesmo id em ambos os lados, já que a versão é
  literalmente uma foto da tabela `atividades` naquele instante).
- Atividade que existe nos DOIS (versão e ao vivo): a linha mostra os dados
  ATUAIS — célula que mudou desde a versão fica com fundo AMARELO; célula
  igual à versão fica sem destaque (mostrar o valor atual e destacar só o
  que mudou é o que o usuário pediu: "apresentar as informações do
  cronograma atual na cor amarela").
- Atividade que só existe no cronograma ATUAL (criada depois da versão):
  linha inteira com fundo AZUL BEM CLARO.
- Atividade que só existe na VERSÃO (excluída depois, no cronograma ao
  vivo): mostra os dados como estavam salvos na versão, sem destaque — não
  tem "atual" pra comparar contra, e o pedido do usuário não cobre esse
  caso; ficou de fora tanto do azul (não é nova) quanto do amarelo (não há
  uma "mesma atividade atual" cujos dados mostrar).
- Ordem das linhas: primeiro todas as atividades da versão (na ordem em que
  foram salvas), depois as atividades novas (só no cronograma atual).

As colunas comparadas são as mesmas — e reaproveitam os mesmos formatadores
— da planilha de "Exportar Cronograma" (cronograma_export.py), MENOS
"Atrasada": é um campo calculado a partir da data de HOJE, não um dado
editado pelo usuário, então "mudou" entre a versão e agora não significa
nada (o calendário sempre anda) — incluí-la na comparação só geraria ruído
amarelo sem relação com uma edição de verdade.

`gerar_planilha_bytes(versao, atividades_atuais, deps_versao_por_atividade,
deps_atuais_por_atividade)`:
- `versao`: o dict devolvido por cronograma_versoes.obter_versao() (usa
  `numero_versao`, `rotulo`, `criado_em` e `dados.atividades`).
- `atividades_atuais`: lista de dicts no formato ATIVIDADE_SELECT (ver
  main.py), já filtrada pelo projeto da versão — SEM outros filtros, a
  comparação é sempre do cronograma inteiro.
- `deps_versao_por_atividade` / `deps_atuais_por_atividade`: dicts
  {atividade_id: [{predecessora_codigo_wbs, tipo, lag_horas}, ...]} — mesmo
  formato usado em cronograma_export.gerar_planilha_bytes, um pra cada lado
  da comparação (a versão já traz suas próprias dependências resolvidas em
  `dados.dependencias`; as atuais vêm de uma consulta fresca no banco).
Nenhuma consulta ao banco acontece aqui, só formatação — os dados já vêm
prontos de quem chama (ver /api/cronograma/versoes/<id>/comparar em
main.py).
"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .cronograma_export import (
    COR_CABECALHO_HEX, _autofit, _depende_de_txt, _fmt_data_br,
    _item_tr_txt, _responsaveis_txt,
)

FILL_AZUL_CLARO = PatternFill(start_color="DCEBFB", end_color="DCEBFB", fill_type="solid")
FILL_AMARELO = PatternFill(start_color="FFF3B0", end_color="FFF3B0", fill_type="solid")

# (cabeçalho, getter(atividade, deps) -> valor da célula) — mesma lista de
# cronograma_export.CABECALHO, sem "Atrasada" (ver docstring acima).
COLUNAS = [
    ("Código", lambda a, deps: a.get("codigo_wbs") or ""),
    ("Atividade", lambda a, deps: a.get("nome") or ""),
    ("Etapa", lambda a, deps: f"{a['etapa_numero']} — {a['etapa_nome']}" if a.get("etapa_nome") else ""),
    ("Frente", lambda a, deps: a.get("frente_nome") or ""),
    ("Status", lambda a, deps: a.get("status") or ""),
    ("Prioridade", lambda a, deps: a.get("prioridade") or ""),
    ("Início prev.", lambda a, deps: _fmt_data_br(a.get("dtini_prev"))),
    ("Fim prev.", lambda a, deps: _fmt_data_br(a.get("dtfim_prev"))),
    ("Esforço prev. (h)", lambda a, deps: float(a["prazo_horas"]) if a.get("prazo_horas") is not None else None),
    ("Início real", lambda a, deps: _fmt_data_br(a.get("dtini_real"))),
    ("Fim real", lambda a, deps: _fmt_data_br(a.get("dtfim_real"))),
    ("Horas realizadas (h)", lambda a, deps: float(a["horas_realizadas"]) if a.get("horas_realizadas") is not None else None),
    ("% concluído", lambda a, deps: a.get("percentual_concluido")),
    ("Master", lambda a, deps: "Sim" if a.get("eh_atividade_master") else ""),
    ("Entregável", lambda a, deps: "Sim" if a.get("eh_entregavel") else ""),
    ("Responsáveis", lambda a, deps: _responsaveis_txt(a.get("responsaveis"))),
    ("Tipo de atividade elementar", lambda a, deps: a.get("tipo_nome") or ""),
    ("Item do TR", lambda a, deps: _item_tr_txt(a)),
    ("Depende de", lambda a, deps: _depende_de_txt(deps)),
    ("Descrição", lambda a, deps: a.get("descricao") or ""),
    ("Objetivo", lambda a, deps: a.get("objetivo") or ""),
    ("Observações", lambda a, deps: a.get("observacoes") or ""),
]
LARGURAS = [12, 36, 22, 20, 16, 11, 13, 13, 14, 13, 13, 15, 12, 8, 10, 28,
            24, 26, 30, 34, 30, 34]


def _fmt_datahora_br(d):
    if not d:
        return "—"
    s = d if isinstance(d, str) else d.isoformat()
    # "YYYY-MM-DDTHH:MM:SS..." (com ou sem timezone/microssegundos) -> "DD/MM/AAAA HH:MM"
    data, _, resto = s.partition("T")
    partes = data.split("-")
    data_br = f"{partes[2]}/{partes[1]}/{partes[0]}" if len(partes) == 3 else data
    hora = resto[:5] if len(resto) >= 5 else ""
    return f"{data_br} {hora}".strip()


def _valores(a, deps_por_atividade):
    deps = (deps_por_atividade or {}).get(a.get("id")) or []
    return [getter(a, deps) for _, getter in COLUNAS]


def gerar_planilha_bytes(versao, atividades_atuais, deps_versao_por_atividade=None, deps_atuais_por_atividade=None):
    atividades_versao = ((versao or {}).get("dados") or {}).get("atividades") or []
    by_id_versao = {a["id"]: a for a in atividades_versao if a.get("id")}
    by_id_atual = {a["id"]: a for a in atividades_atuais if a.get("id")}
    n_colunas = len(COLUNAS)

    wb = Workbook()
    ws = wb.active
    ws.title = "Comparação"

    titulo = (
        f"Comparação — Versão {versao.get('numero_versao')}"
        + (f' "{versao.get("rotulo")}"' if versao.get("rotulo") else "")
        + f" (gerada em {_fmt_datahora_br(versao.get('criado_em'))})"
        + f"  ×  Cronograma atual (gerado em {_fmt_datahora_br(datetime.now())})"
    )
    ws.append([titulo])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_colunas)
    ws.cell(1, 1).font = Font(bold=True, size=12)

    legenda_row = 2
    ws.cell(legenda_row, 1, "Nova desde a versão").fill = FILL_AZUL_CLARO
    ws.cell(legenda_row, 3, "Campo alterado desde a versão (mostra o valor ATUAL)").fill = FILL_AMARELO

    header_row = 3
    for i, (nome, _) in enumerate(COLUNAS, start=1):
        cel = ws.cell(header_row, i, nome)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill(start_color=COR_CABECALHO_HEX, end_color=COR_CABECALHO_HEX, fill_type="solid")

    linha = header_row + 1

    # 1) todas as atividades da VERSÃO, na ordem salva.
    for a_versao in atividades_versao:
        aid = a_versao.get("id")
        a_atual = by_id_atual.get(aid) if aid else None
        if a_atual is not None:
            valores_versao = _valores(a_versao, deps_versao_por_atividade)
            valores_atual = _valores(a_atual, deps_atuais_por_atividade)
            for col_i, valor in enumerate(valores_atual, start=1):
                cel = ws.cell(linha, col_i, valor)
                if valores_versao[col_i - 1] != valor:
                    cel.fill = FILL_AMARELO
        else:
            # excluída do cronograma ao vivo depois da versão — mostra como
            # estava salva, sem destaque (ver docstring do módulo).
            for col_i, valor in enumerate(_valores(a_versao, deps_versao_por_atividade), start=1):
                ws.cell(linha, col_i, valor)
        linha += 1

    # 2) atividades que só existem no cronograma ATUAL (não estavam na versão).
    for a_atual in atividades_atuais:
        aid = a_atual.get("id")
        if aid and aid in by_id_versao:
            continue
        for col_i, valor in enumerate(_valores(a_atual, deps_atuais_por_atividade), start=1):
            cel = ws.cell(linha, col_i, valor)
            cel.fill = FILL_AZUL_CLARO
        linha += 1

    _autofit(ws, LARGURAS)
    ws.freeze_panes = f"A{header_row + 1}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
