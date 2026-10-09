"""
Exportação em Excel do quadro "Rubricas × Situação" — Dashboard Comparação
Folha (152ª rodada, botão de exportar pedido na 153ª).

Diferente dos outros exports deste app (cronograma, quadro de ciclo de
migração, relatório de atividades — todos com colunas FIXAS, uma constante
CABECALHO no topo do arquivo), aqui as colunas não são fixas: vêm dos
dados, são as mesmas situações calculadas por
comparacao_folha.por_rubrica_situacao() pra competência selecionada. Por
isso o cabeçalho é montado dinamicamente a partir de `situacoes`, em vez de
uma constante.

Mesmo espírito de app/migracao_quadro_export.py: uma aba só, cabeçalho
colorido, linha de total geral em negrito ao final — nenhum recálculo
acontece aqui, só formatação do que `por_rubrica_situacao()` já devolveu.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

COR_CABECALHO_HEX = "1A3D5C"
COR_TOTAL_HEX = "E8EDF2"

_CARACTERES_INVALIDOS_ABA = set(':\\/?*[]')


def _titulo_aba(mesano):
    """Nome de aba do Excel não aceita : \\ / ? * [ ] e tem limite de 31
    caracteres — sanitiza pra nunca quebrar wb.save() com um mesano/sufixo
    inesperado."""
    bruto = f"Rubricas x Situacao {mesano}" if mesano else "Rubricas x Situacao"
    limpo = "".join(c for c in bruto if c not in _CARACTERES_INVALIDOS_ABA)
    return (limpo[:31] or "Rubricas x Situacao")


def gerar_planilha_bytes(dados, mesano=None):
    """`dados`: mesmo dict devolvido por comparacao_folha.por_rubrica_situacao()
    — {"situacoes": [...], "totais_situacao": [...], "linhas": [{"rubrica_ergon",
    "rubrica_nome_ergon", "total", "valores": [...]}], "total_geral": N}.
    `mesano`: só usado pro título da aba (nome de arquivo é resolvido por
    quem chama, em main.py)."""
    situacoes = dados.get("situacoes") or []
    totais_situacao = dados.get("totais_situacao") or []
    linhas = dados.get("linhas") or []
    total_geral = dados.get("total_geral") or 0

    wb = Workbook()
    ws = wb.active
    ws.title = _titulo_aba(mesano)

    cabecalho = ["Rubrica Ergon", "Nome"] + situacoes + ["Total"]
    ws.append(cabecalho)
    for col in range(1, len(cabecalho) + 1):
        cel = ws.cell(row=1, column=col)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill(start_color=COR_CABECALHO_HEX, end_color=COR_CABECALHO_HEX, fill_type="solid")
        cel.alignment = Alignment(vertical="center")

    linha = 2
    for l in linhas:
        valores = l.get("valores") or []
        ws.append([
            l.get("rubrica_ergon") or "",
            l.get("rubrica_nome_ergon") or "",
            *[v or 0 for v in valores],
            l.get("total") or 0,
        ])
        linha += 1

    if linhas:
        ws.append(["Total", "", *[v or 0 for v in totais_situacao], total_geral])
        for col in range(1, len(cabecalho) + 1):
            cel = ws.cell(row=linha, column=col)
            cel.font = Font(bold=True)
            cel.fill = PatternFill(start_color=COR_TOTAL_HEX, end_color=COR_TOTAL_HEX, fill_type="solid")

    larguras = [16, 34] + [22] * len(situacoes) + [10]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    # Congela Rubrica Ergon + Nome (colunas A-B) e o cabeçalho (linha 1) —
    # com várias colunas de situação a lista fica larga, e essas duas
    # colunas são a referência de qual linha é qual ao rolar pra direita.
    ws.freeze_panes = "C2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
